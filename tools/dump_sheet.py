"""Ekspor MASTER KARYAWAN jadi array 2D JSON, meniru getDataRange().getValues()."""
import sys, json
from pathlib import Path
from datetime import datetime, date
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

# Path relatif terhadap repo - folder ini sudah dua kali rename/pindah
# (0. EMPLOYEE NS RECORD -> 0. HR DASHBOARD, dan beda root drive), sehingga
# hardcode absolut sebelumnya basi dan script gagal total (ENOENT).
p = Path(__file__).parent.parent / "REF" / "EMPLOYEE DATA.xlsx"
ws = openpyxl.load_workbook(p, data_only=True)["MASTER KARYAWAN"]

def cell(v):
    # GAS mengembalikan Date object; kita tandai supaya harness node bisa rekonstruksi
    if isinstance(v, (datetime, date)):
        return {"__date__": v.isoformat()}
    if v is None:
        return ""
    return v

rows = []
# MASTER KARYAWAN punya 35 kolom (A:AI) - sebelumnya berhenti di 32 (AF) sehingga
# TANGGAL EFEKTIF NON AKTIF, ALASAN KELUAR, dan PENDIDIKAN tidak pernah masuk
# fixture dan harness buta terhadap kolom-kolom itu (audit D1).
for r in range(1, ws.max_row + 1):
    rows.append([cell(ws.cell(row=r, column=c).value) for c in range(1, 36)])

out = Path(__file__).parent / "sheet_values.json"
out.write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")
print(f"rows={len(rows)} cols={len(rows[0])} -> {out}")
